import re
from io import BytesIO

import streamlit as st
from openpyxl import load_workbook


# ======================================================
# PAGE SETTINGS
# ======================================================

st.set_page_config(
    page_title="Model Info Converter",
    page_icon="📏",
    layout="centered"
)

st.title("📏 Model Info Converter")

st.write(
    "Upload an Excel file to convert legacy Model_Info values into the new format."
)


# ======================================================
# CONVERSION FUNCTIONS
# ======================================================

def normalize_size(size):

    size = size.strip()

    # Remove inch symbols
    size = size.replace('"', '')

    # Standardize spacing
    size = re.sub(r"\s*[xX×]\s*", " x ", size)

    size_map = {
        "small": "S",
        "medium": "M",
        "large": "L"
    }

    return size_map.get(size.lower(), size)


def convert_model_info(text):

    if not text:
        return None

    text = str(text).strip()

    # Remove HTML breaks
    text = re.sub(
        r"<br\s*/?>",
        "",
        text,
        flags=re.IGNORECASE
    ).strip()

    # ==========================
    # HEIGHT
    # Supports:
    # 5'9"
    # 5'9
    # 6'
    # 6'0"
    # ==========================

    height_match = re.search(
        r"(\d+)'\s*(\d*)\"?",
        text
    )

    if not height_match:
        return None

    feet = int(height_match.group(1))

    inch_part = height_match.group(2)

    inches = int(inch_part) if inch_part else 0

    cm = round(
        ((feet * 12) + inches) * 2.54
    )

    output = (
        f'Model is {cm}cm '
        f'({feet}\'{inches}")'
    )

    # ==========================
    # WAIST
    # ==========================

    waist_match = re.search(
        r'waist\s*(\d+)"?',
        text,
        flags=re.IGNORECASE
    )

    if not waist_match:

        waist_match = re.search(
            r'(\d+)"\s*waist',
            text,
            flags=re.IGNORECASE
        )

    if waist_match:

        waist = waist_match.group(1)

        output += (
            f', Waist {waist}"'
        )

    # ==========================
    # SIZE
    # ==========================

    size_match = re.search(
        r'wearing(?:\s+a)?\s+size\s+([^.,<]+)',
        text,
        flags=re.IGNORECASE
    )

    if size_match:

        size = normalize_size(
            size_match.group(1)
        )

        output += (
            f', Wearing a Size {size}'
        )

    return output


# ======================================================
# FILE UPLOAD
# ======================================================

uploaded_file = st.file_uploader(
    "Upload Excel File",
    type=["xlsx"]
)

if uploaded_file:

    try:

        wb = load_workbook(uploaded_file)
        ws = wb.active

        headers = []

        for cell in ws[1]:
            if cell.value:
                headers.append(
                    str(cell.value).strip()
                )

        st.success(
            "File loaded successfully!"
        )

        row_count = ws.max_row - 1

        st.write(
            f"Rows found: {row_count}"
        )

        required_columns = [
            "PC-9",
            "PC9",
            "Model_Info"
        ]

        missing = []

        for col in required_columns:

            if col not in headers:
                missing.append(col)

        if missing:

            st.error(
                "Missing required column(s): "
                + ", ".join(missing)
            )

        else:

            st.success(
                "All required columns found."
            )

            if st.button("Convert"):

                header_map = {}

                for cell in ws[1]:

                    if cell.value:

                        header_map[
                            str(cell.value).strip()
                        ] = cell.column

                model_info_col = header_map[
                    "Model_Info"
                ]

                output_col = 4

                ws.cell(
                    row=1,
                    column=output_col,
                    value="Converted_Model_Info"
                )

                processed = 0
                converted = 0
                unmatched = 0

                for row in range(
                    2,
                    ws.max_row + 1
                ):

                    processed += 1

                    source_text = ws.cell(
                        row=row,
                        column=model_info_col
                    ).value

                    result = convert_model_info(
                        source_text
                    )

                    if result:

                        converted += 1

                        ws.cell(
                            row=row,
                            column=output_col,
                            value=result
                        )

                    else:

                        unmatched += 1

                        ws.cell(
                            row=row,
                            column=output_col,
                            value=""
                        )

                output = BytesIO()

                wb.save(output)

                output.seek(0)

                st.success(
                    "Conversion complete!"
                )

                st.write("### Conversion Summary")

                st.write(
                    f"Rows Processed: {processed}"
                )

                st.write(
                    f"Converted: {converted}"
                )

                st.write(
                    f"Unmatched: {unmatched}"
                )

                st.download_button(
                    label="Download Output File",
                    data=output,
                    file_name="output.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

    except Exception as e:

        st.error(str(e))